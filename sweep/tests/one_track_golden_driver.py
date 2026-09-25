"""Test-only driver for sweep/tests/test_one_track_goldens.py.

Always run in a FRESH interpreter, never imported by the suite. `config`
selects its profile at import and `scraper` compiles its scoring tables and
binds the profile's engine stamp at import (config.py:1175-1188,
scraper.py:311-344). So the only faithful way to observe one profile's
behaviour is a new process per profile, exactly as the engine itself runs.

Two modes, each reading a JSON spec on stdin and printing JSON on stdout:

  render  make_profile.render() for every case, in a process whose config is
          NOT overlaid by any profile. That is how Render and the worker
          render. The engine stamp comes from this process's environment.
  run     one rendered profile, loaded by name through the real import path
          (a temporary `profiles` package placed first on sys.path). Reports
          the engine binding, the dry-run plan, each unit's built provider
          input, the combo keys and, when rows are given, scoring, finalize,
          the score_job mutation contract, outputs and exports.

It writes only inside the working directory the caller gives it (a temp
dir). It opens no network connection and reads no .env.
"""
import contextlib
import copy
import csv
import io
import json
import os
import socket
import sys

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIXED_DAY = "2026-01-01"          # combo-key date; the real one is today's date
FIXED_EXPORTED = "2026-01-01 00:00"


def _no_network(*_a, **_k):
    raise RuntimeError("the golden driver must not open a network connection")


socket.socket.connect = _no_network
socket.create_connection = _no_network


def _paths(first=None):
    for path in [p for p in (first,) if p] + [REPO, os.path.join(REPO, "auto-apply")]:
        if path not in sys.path:
            sys.path.insert(0 if path == first else len(sys.path), path)


def render(spec):
    sys.argv = ["render"]
    _paths()
    import make_profile
    import skill_concepts
    return {"engine_version": skill_concepts.engine_version(),
            "sources": {case["name"]: make_profile.render(case["name"], case["derived"],
                                                          case["prefs"])
                        for case in spec["cases"]}}


def _plan(scraper, config):
    """The dry run exactly as --dry-run --json prints it, plus what the JSON
    leaves out: each unit dict and the provider input built from it."""
    from sweep import runs
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        scraper.main()
    dry = json.loads(out.getvalue())
    args = scraper.parse_args()
    units = {}
    for site in scraper.resolve_sites(args):
        plan = scraper.plan_for_site(site, args)
        if plan:
            units[site] = [{"search": unit,
                            "input": scraper.build_input(site,
                                                         scraper.effective_search(site, unit))}
                           for unit in plan]
    return {"dry_run": dry, "units": units,
            "combo_keys": runs.combo_keys(dry, FIXED_DAY),
            "search_count": sum(len(v) for v in dry["sites"].values()),
            "profile_changed": list(config.PROFILE_CHANGED)}


def _strip(rows):
    return [{k: v for k, v in row.items() if k != "_case"} for row in rows]


def _scoring(scraper, rows, name):
    from sweep import exports, logic
    base = _strip(rows)

    # score_job, one row at a time, on its own copy: what it keeps, what it
    # adds to the row, and what it leaves alone.
    per_row = []
    for row in base:
        mine = copy.deepcopy(row)
        before = copy.deepcopy(mine)
        kept = scraper.score_job(mine) is not None
        added = sorted(set(mine) - set(before))
        changed = sorted(k for k in before if mine.get(k) != before[k])
        per_row.append({"url": row["Job URL"], "kept": kept,
                        "added_keys": added, "changed_keys": changed,
                        "added": {k: mine[k] for k in added}})

    # Mutation contract on two controlled rows (index 0 kept, 6 dropped).
    kept_row = copy.deepcopy(base[0])
    scraper.score_job(kept_row)
    first = copy.deepcopy(kept_row)
    scraper.score_job(kept_row)
    dropped_row = copy.deepcopy(base[6])
    dropped_before = copy.deepcopy(dropped_row)
    dropped_result = scraper.score_job(dropped_row)
    memo_row = copy.deepcopy(base[0])
    memo_first = scraper._score_once(memo_row)
    memo_flag = memo_row.get(scraper.SCORED)
    stale_verdict = copy.deepcopy(base[0])
    stale_verdict[scraper.SCORED] = False           # a verdict another profile left
    mutation = {
        "repeat_score_job_identical": kept_row == first,
        "dropped_returns_none": dropped_result is None,
        "dropped_row_unchanged": dropped_row == dropped_before,
        "score_once_keeps": memo_first is not None,
        "score_once_flag": memo_flag,
        "score_once_flag_key": scraper.SCORED,
        "score_once_honours_a_stale_false_flag": scraper._score_once(stale_verdict) is None,
        "stale_flag_row_would_be_kept_by_score_job":
            scraper.score_job(copy.deepcopy(base[0])) is not None,
    }

    # score_and_filter's boundary counts, through its own stage hook.
    stages = []
    eligible, stats = scraper.score_and_filter(
        copy.deepcopy(base), stage=lambda stage, got: stages.append([stage, len(got)]))

    out_rows = scraper.finalize(copy.deepcopy(base))
    last_stats = dict(scraper.LAST_STATS)
    memo_rows = copy.deepcopy(base)
    memo_out = scraper.finalize(memo_rows, memo=True)

    # The engine's own files, exactly as write_outputs writes them.
    csv_path, json_path = os.path.abspath("jobs_golden.csv"), os.path.abspath("jobs_golden.json")
    scraper.write_outputs(out_rows, csv_path, json_path)
    raw_csv = open(csv_path, "rb").read()
    engine_json = json.load(open(json_path, encoding="utf-8"))

    # What the results screen and every export read: the CSV back through
    # DictReader, flagged the way rows_with_applied flags it.
    with open(csv_path, newline="", encoding="utf-8") as fh:
        read_back = list(csv.DictReader(fh))
    for row in read_back:
        row["_key"] = scraper._seen_key(row)
        row["applied"] = False
    shown = logic.shortlist(read_back, 0, "", "", logic.DEFAULT_SORT)
    tagged = exports.rows_for_export(logic.bucket_rows(shown), logic.SECTIONS)
    about = [("Profile", name), ("Exported", FIXED_EXPORTED), ("Listings", len(tagged)),
             ("Minimum score", "no minimum"), ("Source", "All sources"),
             ("Search text", "none"), ("Sorted by", logic.SORTS[logic.DEFAULT_SORT][0])]
    xlsx = _xlsx_cells(exports.as_xlsx(tagged, about=about))

    return {
        "score_job": per_row,
        "mutation": mutation,
        "score_and_filter": {"stages": stages, "stats": stats,
                             "eligible_urls": [r["Job URL"] for r in eligible]},
        "finalize": {"rows": out_rows, "last_stats": last_stats,
                     "memo_output_identical": memo_out == out_rows,
                     "memo_flags": [r.get(scraper.SCORED) for r in memo_rows]},
        "engine_outputs": {"csv_has_bom": raw_csv.startswith(b"\xef\xbb\xbf"),
                           "csv_lines": raw_csv.decode("utf-8-sig").splitlines(),
                           "json_equals_finalize": engine_json == out_rows,
                           "output_columns": list(scraper.OUTPUT_COLUMNS)},
        "exports": {"sections": [[key, [r["apply_url"] for r in rows]]
                                 for key, rows in logic.bucket_rows(shown).items()],
                    "csv_has_bom": exports.as_csv(tagged).startswith(b"\xef\xbb\xbf"),
                    "csv_lines": exports.as_csv(tagged).decode("utf-8-sig").splitlines(),
                    "json": json.loads(exports.as_json(tagged)),
                    "html_lines": exports.as_html(tagged, about=about).decode().splitlines(),
                    "xlsx": xlsx},
    }


def _xlsx_cells(body):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(body))
    out = {"sheets": book.sheetnames}
    for sheet in book.worksheets:
        out[sheet.title] = {
            "cells": [[cell.value for cell in row] for row in sheet.iter_rows()],
            "hyperlinks": sorted([cell.coordinate, cell.hyperlink.target]
                                 for row in sheet.iter_rows() for cell in row
                                 if cell.hyperlink is not None),
            "freeze_panes": sheet.freeze_panes,
            "auto_filter": sheet.auto_filter.ref}
    return out


def run(spec):
    name = spec["name"]
    sys.argv = ["scraper.py", "--profile", name, "--dry-run", "--json"]
    _paths(first=os.getcwd())                 # the temporary profiles package wins
    import config
    import scraper
    import skill_concepts
    effective = skill_concepts.effective()
    result = {"matrix": {"profile_engine": config.PROFILE_ENGINE,
                         "profile_schema": getattr(sys.modules[f"profiles.{name}"],
                                                   "PROFILE_SCHEMA", None),
                         "effective_version": effective["version"],
                         "effective_source": effective["source"],
                         "concept_scoring": skill_concepts.enabled(),
                         "evidence": skill_concepts.evidence_enabled()},
              "plan": _plan(scraper, config)}
    if spec.get("rows"):
        result["scoring"] = _scoring(scraper, spec["rows"], name)
    return result


if __name__ == "__main__":
    mode = sys.argv[1]
    spec = json.load(sys.stdin)
    # NOT sort_keys: dict order is part of several contracts here — the
    # dry-run "sites" order IS the plan order the allocator's prefix walks.
    print(json.dumps({"render": render, "run": run}[mode](spec),
                     ensure_ascii=False, default=str))
